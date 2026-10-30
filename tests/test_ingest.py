import io

import pytest
from docx import Document

from app import ingest


def test_plain_text_and_markdown():
    assert ingest.extract_text("a.txt", "Hej\r\n\r\n\r\n\r\nvärlden".encode("utf-8")) == "Hej\n\nvärlden"
    assert ingest.extract_text("b.md", b"# Title") == "# Title"


def test_cp1252_text_falls_back():
    assert ingest.extract_text("a.txt", "Göteborg".encode("cp1252")) == "Göteborg"


def test_docx_paragraphs_and_tables():
    doc = Document()
    doc.add_paragraph("Erfarenhet")
    table = doc.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "2020"
    table.rows[0].cells[1].text = "Producent"
    buffer = io.BytesIO()
    doc.save(buffer)
    text = ingest.extract_text("cv.docx", buffer.getvalue())
    assert "Erfarenhet" in text and "2020 | Producent" in text


def test_pdf_with_no_text_is_rejected():
    from pypdf import PdfWriter

    writer = PdfWriter()
    writer.add_blank_page(width=200, height=200)
    buffer = io.BytesIO()
    writer.write(buffer)
    with pytest.raises(ingest.IngestError, match="No text"):
        ingest.extract_text("scan.pdf", buffer.getvalue())


def test_unsupported_extension():
    with pytest.raises(ingest.IngestError, match="Unsupported"):
        ingest.extract_text("cv.exe", b"x")


def test_source_crud_and_corpus():
    first = ingest.add_source("Old CV", "file", "text one")
    ingest.add_source("A note", "note", "text two")
    assert [s["name"] for s in ingest.list_sources()] == ["A note", "Old CV"]
    corpus = ingest.build_corpus()
    assert "=== SOURCE: Old CV ===\ntext one" in corpus and "=== SOURCE: A note ===" in corpus
    ingest.update_source(first["id"], "Renamed", "changed")
    assert ingest.get_source(first["id"])["text"] == "changed"
    ingest.delete_source(first["id"])
    assert len(ingest.list_sources()) == 1


def test_source_id_is_validated():
    with pytest.raises(KeyError):
        ingest.get_source("../../etc/passwd")

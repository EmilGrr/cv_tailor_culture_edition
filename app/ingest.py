import io
import json
import re
import uuid
from pathlib import Path

from docx import Document
from pypdf import PdfReader

from .config import get_settings

SUPPORTED = {".pdf", ".docx", ".txt", ".md"}


class IngestError(ValueError):
    pass


def extract_text(filename: str, data: bytes) -> str:
    ext = Path(filename).suffix.lower()
    if ext not in SUPPORTED:
        raise IngestError(f"Unsupported file type '{ext}'. Use PDF, DOCX, TXT or MD.")
    if ext == ".pdf":
        text = _pdf_text(data)
    elif ext == ".docx":
        text = _docx_text(data)
    else:
        text = _plain_text(data)
    text = _normalise(text)
    if not text:
        raise IngestError(
            f"No text could be extracted from '{filename}'. Scanned documents are not supported."
        )
    return text


def _pdf_text(data: bytes) -> str:
    try:
        reader = PdfReader(io.BytesIO(data))
        return "\n".join((page.extract_text() or "") for page in reader.pages)
    except Exception as exc:
        raise IngestError(f"Could not read PDF: {exc}") from exc


def _docx_text(data: bytes) -> str:
    try:
        doc = Document(io.BytesIO(data))
    except Exception as exc:
        raise IngestError(f"Could not read DOCX: {exc}") from exc
    lines = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                lines.append(" | ".join(cells))
    return "\n".join(lines)


def _plain_text(data: bytes) -> str:
    for encoding in ("utf-8-sig", "cp1252"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    return data.decode("latin-1")


def _normalise(text: str) -> str:
    text = text.replace("\r\n", "\n").replace("\r", "\n").replace("\x00", "")
    text = re.sub(r"[ \t]+\n", "\n", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def _path(source_id: str) -> Path:
    if not re.fullmatch(r"[0-9a-f]{8}", source_id):
        raise KeyError(source_id)
    return get_settings().sources_dir / f"{source_id}.json"


def _write(source: dict) -> None:
    _path(source["id"]).write_text(json.dumps(source, ensure_ascii=False, indent=2), encoding="utf-8")


def add_source(name: str, kind: str, text: str) -> dict:
    get_settings().sources_dir.mkdir(parents=True, exist_ok=True)
    source = {"id": uuid.uuid4().hex[:8], "name": name.strip() or "Untitled", "kind": kind, "text": text}
    _write(source)
    return source


def update_source(source_id: str, name: str, text: str) -> dict:
    source = get_source(source_id)
    source["name"] = name.strip() or source["name"]
    source["text"] = text
    _write(source)
    return source


def get_source(source_id: str) -> dict:
    path = _path(source_id)
    if not path.exists():
        raise KeyError(source_id)
    return json.loads(path.read_text(encoding="utf-8"))


def list_sources() -> list[dict]:
    sources_dir = get_settings().sources_dir
    if not sources_dir.exists():
        return []
    sources = [json.loads(p.read_text(encoding="utf-8")) for p in sources_dir.glob("*.json")]
    return sorted(sources, key=lambda s: s["name"].lower())


def delete_source(source_id: str) -> None:
    path = _path(source_id)
    if not path.exists():
        raise KeyError(source_id)
    path.unlink()


def build_corpus() -> str:
    parts = [f"=== SOURCE: {s['name']} ===\n{s['text']}" for s in list_sources()]
    return "\n\n".join(parts)
